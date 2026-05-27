from fastapi import FastAPI, File, UploadFile, Form
from fastapi.responses import HTMLResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from groq import Groq
from dotenv import load_dotenv
import os
import base64
import io
import docx
import openpyxl

load_dotenv()

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

client = Groq(api_key=os.getenv("GROQ_API_KEY"))

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
INDEX_HTML_PATH = os.path.join(BASE_DIR, "index.html")

class Message(BaseModel):
    message: str
    system: str = "You are a helpful assistant."

def extract_text_from_docx_bytes(file_bytes: bytes) -> str:
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        return "\n".join([para.text for para in doc.paragraphs])
    except Exception as e:
        return f"[Error parsing Word file: {str(e)}]"

def extract_text_from_xlsx_bytes(file_bytes: bytes) -> str:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        text_content = []
        for sheet in wb.worksheets:
            text_content.append(f"--- Sheet: {sheet.title} ---")
            for row in sheet.iter_rows(values_only=True):
                if any(row):
                    text_content.append(" | ".join([str(v) if v is not None else "" for v in row]))
        return "\n".join(text_content)
    except Exception as e:
        return f"[Error parsing Excel file: {str(e)}]"

def extract_text_from_pdf_bytes(file_bytes: bytes) -> str:
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        return "\n".join([page.extract_text() or "" for page in reader.pages])
    except Exception as e:
        return f"[Error parsing PDF file: {str(e)}]"

@app.get("/", response_class=HTMLResponse)
async def read_root():
    with open(INDEX_HTML_PATH, "r", encoding="utf-8") as f:
        return f.read()

@app.post("/chat")
async def chat(msg: Message):
    try:
        response = client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            messages=[
                {"role": "system", "content": msg.system},
                {"role": "user", "content": msg.message}
            ]
        )
        return {"reply": response.choices[0].message.content}
    except Exception as e:
        return {"reply": f"Error: {str(e)}"}

@app.post("/upload")
async def upload_file(
    file: UploadFile = File(...),
    message: str = Form(""),
    system: str = Form("You are a helpful assistant.")
):
    try:
        if not file.filename:
            return JSONResponse(status_code=400, content={"reply": "No file was selected."})

        file_extension = os.path.splitext(file.filename)[1].lower()
        file_bytes = await file.read()

        # ── IMAGE ──
        if file_extension in [".png", ".jpg", ".jpeg", ".webp"]:
            mime_type = "image/jpeg"
            if file_extension == ".png":
                mime_type = "image/png"
            elif file_extension == ".webp":
                mime_type = "image/webp"

            encoded_image = base64.b64encode(file_bytes).decode("utf-8")
            prompt = message.strip() or "Please describe this image in detail."

            response = client.chat.completions.create(
                model="llama-3.2-11b-vision-preview",
                messages=[
                    {"role": "system", "content": system},
                    {
                        "role": "user",
                        "content": [
                            {"type": "text", "text": prompt},
                            {
                                "type": "image_url",
                                "image_url": {"url": f"data:{mime_type};base64,{encoded_image}"}
                            }
                        ]
                    }
                ]
            )
            reply = response.choices[0].message.content

        # ── AUDIO ──
        elif file_extension in [".wav", ".mp3", ".m4a", ".webm", ".ogg"]:
            transcription = client.audio.transcriptions.create(
                model="whisper-large-v3",
                file=(file.filename, file_bytes)
            )
            extracted_text = transcription.text
            prompt = message.strip() or "Please analyze this voice message and respond."
            full_prompt = f'User Voice Message (Transcribed): "{extracted_text}"\n\nUser Query: {prompt}'

            response = client.chat.completions.create(
                model="llama-3.3-70b-versatile",
                messages=[
                    {"role": "system", "content": system},
                    {"role": "user", "content": full_prompt}
                ]
            )
            reply = response.choices[0].message.content
            return {"reply": reply, "filename": file.filename, "type": file_extension, "transcribed": extracted_text}

        # ── DOCUMENTS ──
        else:
            if file_extension == ".docx":
                extracted_text = extract_text_from_docx_bytes(file_bytes)
            elif file_extension in [".xlsx", ".xls"]:
                extracted_text = extract_text_from_xlsx_bytes(file_bytes)
            elif file_extension == ".txt":
                extracted_text = file_bytes.decode("utf-8", errors="ignore")
            elif file_extension == ".pdf":
                extracted_text = extract_text_from_pdf_bytes(file_bytes)
            else:
                extracted_text = f"[Unsupported format '{file_extension}'. Supported: images, audio, PDF, DOCX, XLSX, TXT]"

            prompt = message.strip() or "Please analyze this file and summarize its content."
            full_prompt = f"User Query: {prompt}\n\n[File: '{file.filename}']:\n{extracted_text}"

            response = client.chat.completions.create(
                model="llama-3.3-70b-versatile",
                messages=[
                    {"role": "system", "content": system},
                    {"role": "user", "content": full_prompt}
                ]
            )
            reply = response.choices[0].message.content

        return {"reply": reply, "filename": file.filename, "type": file_extension}

    except Exception as e:
        import traceback
        traceback.print_exc()
        return JSONResponse(status_code=500, content={"reply": f"Error processing file: {str(e)}"})
